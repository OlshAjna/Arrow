from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
FONT = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(90, 98, 108)
HEADER_FILL = "D9EAF7"
ALT_FILL = "F7FAFC"
BORDER_COLOR = "D9D9D9"


def set_run_font(run, size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def configure_paragraph_style(style, size, bold=False, before=0, after=6):
    style.font.name = FONT
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.2


def set_section_portrait(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)


def set_section_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)


def remove_paragraph_borders(style):
    p_pr = style.element.get_or_add_pPr()
    for border in p_pr.findall(qn("w:pBdr")):
        p_pr.remove(border)


def configure_document(doc, title, subject):
    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "Arrow 课程设计"
    props.keywords = "Cocos Creator, 方向箭头, 系统设计, 课程设计"

    styles = doc.styles
    configure_paragraph_style(styles["Normal"], 10.5, after=6)
    configure_paragraph_style(styles["Title"], 23, bold=True, before=10, after=7)
    configure_paragraph_style(styles["Heading 1"], 15, bold=True, before=16, after=7)
    configure_paragraph_style(styles["Heading 2"], 12.5, bold=True, before=11, after=5)
    configure_paragraph_style(styles["List Bullet"], 10.5, after=4)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.color.rgb = BLACK
        remove_paragraph_borders(styles[style_name])
    set_section_portrait(doc.sections[0])


def add_paragraph(doc, text, bold_prefix=None, align=None, space_after=6):
    p = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        lead = p.add_run(bold_prefix)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.26)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix and text.startswith(bold_prefix):
        lead = p.add_run(bold_prefix)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_numbered_item(doc, number, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    p.paragraph_format.space_after = Pt(4)
    lead = p.add_run(f"{number}. ")
    set_run_font(lead, bold=True)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER_COLOR)


def set_cell_margins(table, top=95, start=105, bottom=95, end=105):
    tbl_pr = table._tbl.tblPr
    cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = cell_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            cell_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths, font_size=9.5, compact=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if compact:
        set_cell_margins(table, top=60, start=90, bottom=60, end=90)
    else:
        set_cell_margins(table)

    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=font_size, bold=True)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, ALT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)

    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_metadata(doc, rows):
    return add_table(
        doc,
        ["项目", "内容", "项目", "内容"],
        rows,
        [0.72, 2.63, 0.72, 2.63],
        font_size=9.2,
    )


def add_figure(doc, path, width_inches, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    set_run_font(run, size=9, color=MUTED, italic=True)
    return p


def build_technical_plan():
    doc = Document()
    configure_document(doc, "技术可行性评估与技术方案", "方向箭头解谜小游戏")

    title = doc.add_paragraph("技术可行性评估与技术方案", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("基于 Cocos Creator 的方向箭头解谜小游戏")
    set_run_font(run, size=11.5, color=MUTED)

    add_metadata(
        doc,
        [
            ["项目", "方向箭头解谜小游戏", "平台", "Cocos Creator 3.x"],
            ["文档用途", "核心玩法可行性评估", "版本", "1.0"],
        ],
    )

    add_heading(doc, "1 项目范围与结论", 1)
    add_paragraph(
        doc,
        "本项目是一款基于固定网格的方向箭头解谜小游戏。玩家点击箭头后，箭头沿自身方向尝试移动；路径畅通时滑出屏幕并消除，路径被墙体或其他箭头阻挡时退回原位。系统还包含生命值、重置、提示、关卡进度和设置保存。",
    )
    add_paragraph(
        doc,
        "计划采用 Cocos Creator 3.x、TypeScript 和微信小游戏平台。综合来看，核心玩法在课程项目规模内技术可行，主要工作集中在网格碰撞判定、箭头状态管理和关卡数据组织，不需要复杂的物理系统或服务器能力。",
    )
    add_table(
        doc,
        ["评估项", "结论", "说明"],
        [
            ["元素绘制", "可行", "箭头、折线、墙体和文字均可由 Cocos 2D 节点或 Graphics 完成"],
            ["点击检测", "可行", "屏幕坐标可以转换为网格坐标，再查询对应箭头"],
            ["移动动画", "可行", "可使用 Cocos Tween 实现直线滑出和退回"],
            ["阻挡判定", "可行", "沿蛇头方向逐格检查墙体和其他箭头即可"],
            ["胜负与道具", "可行", "由游戏状态、剩余数量和提示扫描完成"],
            ["本地存储", "可行", "使用微信小程序本地存储接口保存单个 JSON 存档"],
        ],
        [1.25, 0.75, 4.55],
    )

    add_heading(doc, "2 核心技术方案", 1)
    add_heading(doc, "2.1 元素绘制", 2)
    add_paragraph(
        doc,
        "棋盘采用从左上角开始的逻辑网格坐标。每个箭头作为一个独立节点，箭头身体使用折线坐标绘制，箭头头部使用图形或预制体标记方向。墙体按照关卡配置生成固定节点。",
    )
    add_paragraph(
        doc,
        "表现层按背景、墙体、箭头、提示与弹窗的顺序绘制。箭头之间的前后关系由 drawOrder 控制，并映射为节点的显示顺序。逻辑坐标与屏幕坐标分开保存，屏幕尺寸变化时只需要重新换算显示位置。",
    )

    add_heading(doc, "2.2 点击与移动", 2)
    add_paragraph(
        doc,
        "玩家触控屏幕后，UIManager 将点击位置转换为网格坐标，再查询该位置的箭头。GameManager 收到点击对象后，调用 CollisionDetector 判断移动是否合法。",
    )
    add_paragraph(
        doc,
        "路径畅通时，播放沿箭头实际移动方向的滑出动画；动画结束后从当前关卡中删除该箭头并检查通关。路径受阻时，播放短距离前进后返回的动画，逻辑位置保持不变。",
    )

    add_heading(doc, "2.3 阻挡检测", 2)
    add_paragraph(
        doc,
        "碰撞检测采用手工网格射线法，不依赖物理引擎。检测过程从箭头头部的下一个格子开始，沿实际移动方向一直检查到边界；只要遇到墙体或其他箭头，就判定为受阻。",
    )
    add_paragraph(
        doc,
        "普通箭头使用显示方向，反转箭头使用显示方向的反方向。检测时必须排除当前箭头自身。对于课程项目采用的棋盘规模，该算法复杂度较低，性能压力很小。",
    )

    add_heading(doc, "2.4 状态与流程", 2)
    add_paragraph(
        doc,
        "GameManager 统一维护当前关卡、生命值和游戏状态。失败时可以按规则扣除生命并重置当前关卡；全部可清除箭头消失后进入通关状态。",
    )
    add_paragraph(
        doc,
        "LevelManager 负责创建当前关卡的箭头和墙体，保存初始配置以便重置，并向 StorageManager 请求读取或写入进度。提示功能扫描当前所有可移动箭头，并高亮第一个可释放对象。",
    )

    add_heading(doc, "2.5 数据存储", 2)
    add_paragraph(
        doc,
        "本地存储只保存可序列化为 JSON 的数据，不保存 Cocos 节点、组件实例、Map 或当前动画状态。存档使用一个固定键和一个根对象，内容包括格式版本、更新时间、关卡进度和系统设置。",
    )

    add_heading(doc, "3 技术架构", 1)
    add_table(
        doc,
        ["模块", "主要职责"],
        [
            ["GameManager", "全局流程、生命值、游戏状态、胜负判定和模块调度"],
            ["LevelManager", "关卡创建、重置、解锁和进度更新"],
            ["ArrowBase 及子类", "保存箭头属性，并通过多态提供移动方向和行为差异"],
            ["Block", "保存墙体占据格子，参与路径阻挡"],
            ["CollisionDetector", "静态实现网格路径检测"],
            ["UIManager", "绘制、动画、点击转换、提示和界面切换"],
            ["StorageManager", "封装微信小程序本地存储"],
        ],
        [1.55, 5.0],
    )
    add_paragraph(
        doc,
        "实现时使用 Cocos Creator 3.x 与 TypeScript。ArrowBase 和三个子类可以通过继承与重写方法体现多态；普通的数值、数组和 JSON 对象即可覆盖当前功能，不需要额外数据库。",
    )

    add_heading(doc, "4 实施步骤", 1)
    steps = [
        "建立网格坐标、方向和基础数据结构。",
        "完成箭头、墙体和棋盘的基础绘制。",
        "完成点击选中、路径检测和移动动画。",
        "接入生命值、重置、提示和胜负判定。",
        "完成关卡进度与设置的本地保存。",
        "使用固定测试关卡验证普通、固定和反转箭头。",
    ]
    for index, step in enumerate(steps, start=1):
        add_numbered_item(doc, index, step)

    doc.add_page_break()
    add_heading(doc, "5 主要风险与处理", 1)
    add_table(
        doc,
        ["风险", "影响", "处理方式"],
        [
            ["固定箭头语义不明确", "可能导致关卡无法通关", "暂作为不可移动障碍，不计入可清除箭头数量"],
            ["反转箭头方向易混淆", "动画与判定方向可能不一致", "统一由 getMoveDirection 返回实际移动方向"],
            ["坐标混用", "碰撞与显示位置不一致", "逻辑层只使用网格坐标，表现层负责坐标换算"],
            ["存档结构升级", "旧数据无法直接读取", "加入 schemaVersion 并补充默认值"],
            ["动画与规则不同步", "提前删除或重复点击", "移动期间锁定箭头，动画结束后再更新关卡"],
        ],
        [1.55, 1.85, 3.15],
    )

    add_heading(doc, "6 结论", 1)
    add_paragraph(
        doc,
        "该方案可以在单机、小规模关卡条件下完成。技术重点不是渲染或平台接口，而是保证路径检测、动画结束时机和关卡状态三者一致。先完成固定关卡和基础箭头，再逐步加入特殊箭头与存储功能，是风险较低的实现顺序。",
    )

    output = DOCS_DIR / "技术可行性评估与技术方案.docx"
    doc.save(output)
    return output


def build_system_design():
    doc = Document()
    configure_document(doc, "系统设计总说明书", "方向箭头解谜小游戏")

    title = doc.add_paragraph("系统设计总说明书", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("方向箭头解谜小游戏")
    set_run_font(run, size=11.5, color=MUTED)

    add_metadata(
        doc,
        [
            ["项目", "方向箭头解谜小游戏", "平台", "Cocos Creator 3.x"],
            ["文档用途", "系统设计说明", "版本", "1.0"],
        ],
    )

    add_heading(doc, "1 设计概述", 1)
    add_paragraph(
        doc,
        "本说明书描述方向箭头解谜小游戏的总体架构、核心类设计和存储数据模型。项目面向 Cocos Creator 3.x 与微信小游戏平台，目标是在单机环境中完成固定关卡加载、箭头移动、阻挡判定、胜负结算、重置、提示和进度保存。",
    )
    add_paragraph(
        doc,
        "系统采用游戏流程、关卡管理、领域对象、表现层和存储层相互分工的结构。核心规则由模型和管理器负责，界面只负责输入、绘制和动画，避免场景节点直接决定游戏结果。",
    )

    add_heading(doc, "2 总体架构", 1)
    add_table(
        doc,
        ["层次", "主要类型", "职责"],
        [
            ["应用层", "GameManager", "调度流程、维护生命与状态、处理胜负"],
            ["关卡层", "LevelManager", "创建、重置和保存关卡，管理解锁进度"],
            ["领域层", "ArrowBase 及子类、Block", "描述可交互对象以及普通、固定、反转行为"],
            ["服务层", "CollisionDetector、StorageManager", "提供路径检测和持久化服务"],
            ["表现层", "UIManager", "绘制、点击、动画、提示和页面切换"],
        ],
        [1.02, 2.08, 3.5],
    )
    add_paragraph(
        doc,
        "GameManager 是全局协调者，并以单例形式提供访问入口。它持有 LevelManager、UIManager 和 StorageManager，通过静态工具 CollisionDetector 完成移动合法性判定。该结构能够满足课程项目中管理者、抽象类、继承、多态和工具类的设计展示要求。",
    )

    add_heading(doc, "3 功能模块", 1)
    add_paragraph(
        doc,
        "关卡模块负责加载创作者指定的固定关卡，生成箭头和墙体，并在重置时恢复初始状态。箭头模块保存位置、方向、类型和显示顺序，并通过子类处理普通、固定和反转行为。规则模块负责判断路径是否畅通、扣除生命、判断胜负以及生成提示。",
    )
    add_paragraph(
        doc,
        "界面模块将点击位置转换为游戏对象，显示棋盘和箭头，并播放滑出、退回及高亮动画。存储模块保存关卡进度和系统设置，不保存运行时节点或动画中间状态。",
    )

    add_heading(doc, "4 核心类设计", 1)
    add_paragraph(
        doc,
        "ArrowBase 是抽象基类，保存箭头头部、身体折线、移动方向、类型和显示顺序，并声明移动、消除、重置和获取实际移动方向的通用接口。三个子类通过重写实现行为差异。",
    )
    add_bullet(doc, "NormalArrow：按显示方向移动。")
    add_bullet(doc, "FixedArrow：不可移动，作为固定障碍处理。")
    add_bullet(doc, "ReverseArrow：沿显示方向的反方向移动。")
    add_paragraph(
        doc,
        "LevelManager 持有当前关卡中的箭头和墙体，负责对象实例化与重置。CollisionDetector 接收箭头、箭头集合和墙体集合作为参数，不保存关卡状态。GameManager 根据检测结果协调动画、生命值和胜负判断。",
    )
    add_table(
        doc,
        ["类型", "关键职责"],
        [
            ["GameManager", "单例入口、流程调度、状态管理、胜负判定"],
            ["LevelManager", "关卡加载、元素创建、重置、解锁和进度"],
            ["ArrowBase 及子类", "公共属性、统一接口，以及普通、固定和反转行为"],
            ["Block", "保存墙体格子并参与阻挡检测"],
            ["CollisionDetector", "静态路径检测和阻挡物查询"],
            ["UIManager", "绘制、点击、动画、弹窗和 HUD"],
            ["StorageManager", "读取、写入和删除本地存档"],
        ],
        [1.55, 5.0],
        font_size=9.0,
        compact=True,
    )

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_landscape(landscape)
    add_heading(doc, "核心类图", 2)
    add_figure(
        doc,
        ROOT / "mermaid-diagram.png",
        9.35,
        "图 1 核心类与关系。图中的 int、bool 和 Object 为 UML 简写，TypeScript 实现时使用对应具体类型。",
    )

    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_portrait(portrait)

    add_heading(doc, "5 核心交互流程", 1)
    add_paragraph(
        doc,
        "玩家点击箭头后，UIManager 返回命中对象，GameManager 调用 CollisionDetector.canMove()。如果路径畅通，箭头进入移动状态，播放滑出动画；动画完成后由 LevelManager 移除对象，GameManager 更新剩余数量和界面。",
    )
    add_paragraph(
        doc,
        "如果路径受阻，UIManager 播放前进后退动画，逻辑坐标不变。若关卡规则判定本次操作失败，则由 GameManager 扣除生命并判断是否需要显示失败弹窗。",
    )
    add_paragraph(
        doc,
        "重置操作由 GameManager 调用 LevelManager.resetLevel()，根据初始关卡配置重新生成对象。提示操作扫描当前可移动箭头，仅高亮一个可立即释放对象，不改变关卡状态。",
    )
    add_paragraph(
        doc,
        "全部可清除箭头消失后，游戏进入通关状态，LevelManager 解锁下一关并通过 StorageManager 保存进度。当前版本暂将 FixedArrow 视为不可清除的固定障碍，不参与剩余可清除数量统计。",
    )

    add_heading(doc, "6 数据模型", 1)
    add_paragraph(
        doc,
        "存储采用一个固定键和一个存档根对象，根对象包含格式版本、更新时间、关卡进度和系统设置。关卡进度只记录已经解锁或完成的关卡；锁定状态由 maxUnlockedLevel 推导。数据对象及字段说明如下。",
    )

    add_table(
        doc,
        ["数据对象", "字段", "说明"],
        [
            ["PLAYER_SAVE_DATA", "schemaVersion、updatedAt", "存档版本和最后更新时间"],
            ["PROGRESS_DATA", "maxUnlockedLevel", "当前最高已解锁关卡"],
            ["LEVEL_PROGRESS", "levelId、status、completedAt", "单关状态和完成时间"],
            ["SETTINGS_DATA", "音乐、音效、振动开关", "系统设置"],
        ],
        [1.65, 2.2, 2.8],
        font_size=9.2,
    )
    add_figure(doc, ROOT / "data.png", 4.25, "图 2 本地存储数据模型。")
    add_paragraph(
        doc,
        "存档内容使用普通 JSON 对象和数组，不保存 ArrowBase、Block、Cocos Node、Map 或 Set。读取存档时先检查 schemaVersion，字段缺失时补充默认值。写入时生成完整根对象并一次性保存，减少不同字段不一致的情况。",
    )

    add_heading(doc, "6.1 存储原则", 2)
    add_bullet(doc, "使用一个固定键保存完整存档根对象，避免多个键之间状态不一致。")
    add_bullet(doc, "通过 schemaVersion 识别旧格式，并在读取时补充缺失字段或执行迁移。")
    add_bullet(doc, "关卡箭头和墙体由 LevelManager 在进入关卡时重新创建，不写入本地存档。")

    doc.add_page_break()
    add_heading(doc, "7 非功能要求与验收", 1)
    add_paragraph(
        doc,
        "系统应在常见手机屏幕比例下保持棋盘居中和点击区域可辨认。单个关卡规模较小时，网格路径检测可以直接同步执行。动画期间应禁止同一箭头重复触发，并在动画结束后再更新关卡集合。",
    )
    add_paragraph(doc, "课程版本可按以下条件验收：", space_after=3)
    checks = [
        "普通、固定和反转箭头能够按各自规则表现。",
        "墙体和其他箭头能够阻止移动，受阻对象能够回到原位。",
        "全部可清除箭头消除后显示通关，生命耗尽后显示失败。",
        "重置能够恢复关卡初始状态，提示能够高亮可释放箭头。",
        "退出并重新进入游戏后，关卡进度和设置可以恢复。",
    ]
    for index, check in enumerate(checks, start=1):
        add_numbered_item(doc, index, check)

    output = DOCS_DIR / "系统设计总说明书.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    outputs = [build_technical_plan(), build_system_design()]
    for output in outputs:
        print(output)
